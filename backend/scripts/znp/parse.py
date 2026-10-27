"""Разбор «Заказ на производство №N-26.xlsx» (ТЗ docs/roles/tz/ZNP-ORDERS.md).
Запуск из папки «для анализа»:  python -I ../backend/scripts/znp/parse.py
Читает  «Заказы на производство/<N-26>/Заказ на производство №<N-26>.xlsx»,
пишет   znp-parsed.json (шапка + строки), znp-img.json + znp-img/ (картинки по строкам).
"""
import glob, json, os, posixpath, re, zipfile
import openpyxl

parsed, imgs = {}, {}
os.makedirs('znp-img', exist_ok=True)
for f in glob.glob('Заказы на производство/*/Заказ на производство №*.xlsx'):
    z = os.path.basename(os.path.dirname(f))
    ws = openpyxl.load_workbook(f, data_only=True).worksheets[0]
    head, rows, hdr = {}, [], None
    for r in ws.iter_rows(values_only=True):
        s = [str(x).strip() if x is not None else '' for x in r]
        if hdr is None:
            if s and s[0].startswith('№ п/п'):
                hdr = s
                continue
            ne = [x for x in s if x]
            if len(ne) >= 2:
                head[ne[0].rstrip(':')] = ne[1]
        elif s and s[0] and re.match(r'^\d', s[0]):
            rows.append(s)
    parsed[z] = {'head': head, 'rows': rows, 'hdr': hdr}
    # картинки: лист1 → drawing → якорь (строка) → media
    zf = zipfile.ZipFile(f)
    out = []
    srel = 'xl/worksheets/_rels/sheet1.xml.rels'
    if srel in zf.namelist():
        for d in re.findall(r'Target="([^"]+drawing[^"]+)"', zf.read(srel).decode()):
            dp = posixpath.normpath(posixpath.join('xl/worksheets', d))
            x = zf.read(dp).decode()
            rels = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', zf.read(posixpath.join(posixpath.dirname(dp), '_rels', posixpath.basename(dp) + '.rels')).decode()))
            for a in re.findall(r'<xdr:(?:twoCellAnchor|oneCellAnchor)[^>]*>(.*?)</xdr:(?:twoCellAnchor|oneCellAnchor)>', x, re.S):
                row = int(re.search(r'<xdr:from>.*?<xdr:row>(\d+)</xdr:row>', a, re.S).group(1)) + 1
                e = re.search(r'r:embed="([^"]+)"', a)
                if not e:
                    continue
                mp = posixpath.normpath(posixpath.join(posixpath.dirname(dp), rels[e.group(1)]))
                data = zf.read(mp)
                fn = 'znp-img/%s_r%d_%s' % (z, row, posixpath.basename(mp))
                open(fn, 'wb').write(data)
                v = ws.cell(row, 1).value
                out.append({'row': row, 'pos': str(v).strip() if v is not None else None, 'file': fn, 'size': len(data)})
    imgs[z] = out
json.dump(parsed, open('znp-parsed.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1, default=str)
json.dump(imgs, open('znp-img.json', 'w', encoding='utf8'), ensure_ascii=False, indent=1)
print('ЗНП:', len(parsed), 'строк:', sum(len(p['rows']) for p in parsed.values()), 'картинок:', sum(len(v) for v in imgs.values()))
